from openpyxl import load_workbook


# Load Workbook
workbook = load_workbook("TestData.xlsx")


# Select Worksheet
sheet = workbook["Sheet1"]


# Read Data
username = sheet.cell(2, 1).value
password = sheet.cell(2, 2).value

print("Username:", username)
print("Password:", password)


# Write Data
sheet.cell(2, 3).value = "Passed"


# Save Workbook
workbook.save("TestData.xlsx")


# Total Rows
print("Total Rows:", sheet.max_row)


# Total Columns
print("Total Columns:", sheet.max_column)


















